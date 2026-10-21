from __future__ import annotations

import csv
import os
import re
import tempfile
import textwrap
from dataclasses import dataclass
from datetime import datetime, timezone
from math import isfinite
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


TEAL = "#0f7c7a"
NAVY = "#142d43"
PALE = "#eef8fc"
BLUE = "#2864e8"
REPORT_FIELDS = (
    "recordType", "label", "reference", "source", "subscriptionId", "subscriptionName", "currency", "term",
    "hourlyCommitment", "baselineCost", "modeledCost", "savingsAmount",
    "savingsPercent", "coveragePercent", "utilizationPercent", "observationHours",
    "unusedCommitment", "overageCost", "remainingPaygCost", "projectedTermSavings",
    "computeHours", "periodHours", "paygRate", "discountPercent", "commitmentCost",
    "breakEvenHours", "breakEvenRuntimePercent", "runtimeMarginHours",
    "projectedTermBaselineCost", "projectedTermModeledCost",
    "computeType", "computeSku", "benefitType", "benefitSku", "benefitCost",
    "licenseType", "computeInventory",
    "impact", "notes",
)
FIELD_LABELS = {
    "computeType": "Compute type", "computeSku": "Compute SKU",
    "benefitType": "Benefit type", "benefitSku": "Benefit / plan SKU (not workload SKU)",
    "benefitCost": "Benefit cost (report period)",
    "licenseType": "Reported license type (not verified benefit savings)",
    "computeInventory": "Current subscription compute inventory (not attributed benefit coverage)",
    "computeHours": "Compute running hours", "periodHours": "Elapsed period hours",
    "paygRate": "Assumed PAYG rate per running hour", "discountPercent": "Assumed discount (%)",
    "commitmentCost": "Commitment cost (whole period)", "breakEvenHours": "Break-even running hours",
    "breakEvenRuntimePercent": "Break-even runtime (%)", "runtimeMarginHours": "Runtime margin above break-even (hours)",
    "projectedTermBaselineCost": "Projected full-term PAYG cost",
    "projectedTermModeledCost": "Projected full-term commitment plus uncovered PAYG cost",
    "subscriptionId": "Subscription ID", "subscriptionName": "Subscription name",
    "recordType": "Record type", "label": "Recommendation / option", "reference": "Evidence reference",
    "source": "Source", "currency": "Currency", "term": "Term",
    "hourlyCommitment": "Hourly commitment", "baselineCost": "Baseline cost (report period)",
    "modeledCost": "Modeled cost (report period)", "savingsAmount": "Savings amount (report period)",
    "savingsPercent": "Savings (%)", "coveragePercent": "Coverage (%)", "utilizationPercent": "Utilization (%)",
    "observationHours": "Report period hours", "unusedCommitment": "Unused commitment cost",
    "overageCost": "Overage cost", "remainingPaygCost": "Remaining PAYG cost",
    "projectedTermSavings": "Projected full-term savings (offline model only)", "impact": "Advisor impact",
    "notes": "Assumptions, warnings and details",
}
_SECRETS = re.compile(
    r"(?i)\b(?:Bearer\s+\S+|"
    r"(?:client_secret|clientsecret|password|accountkey|sig|access_token|refresh_token)"
    r"\s*[=:]\s*[^\s;,]+)"
    r"|\beyJ[\w-]+\.[\w-]+\.[\w-]+\b"
    r"|\b(?:gh[pousr]_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,})\b"
)


def clean_text(value: Any) -> str:
    if value is None:
        return "Not supplied"
    text = _SECRETS.sub("[REDACTED]", str(value))
    return "".join(character for character in text if ord(character) >= 32 or character in "\n\t")


def redacted_result(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: "[REDACTED]" if re.search(
                r"(?i)authorization|cookie|password|secret|token|connectionstring|accountkey", key
            ) else redacted_result(item)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [redacted_result(item) for item in value]
    if isinstance(value, str):
        return clean_text(value)
    return value


def number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value):
        return float(value)
    return None


def display(value: Any) -> str:
    if value is None:
        return "Not supplied"
    if isinstance(value, float):
        return f"{value:,.2f}"
    return clean_text(value)


@dataclass(frozen=True)
class SavingsReport:
    title: str
    mode: str
    collected_at: str
    scope: str
    context: tuple[str, ...]
    coverage: tuple[str, ...]
    notices: tuple[str, ...]
    records: tuple[dict[str, Any], ...]
    metrics: tuple[tuple[str, str], ...]


def collection_warnings(result: dict[str, Any]) -> tuple[str, ...]:
    labels = {
        "advisor": "Advisor recommendations", "savingsPlan": "Savings Plan recommendations",
        "virtualMachines": "Virtual machine inventory",
        "virtualMachineScaleSets": "VM scale set inventory", "appServicePlans": "App Service plan inventory",
    }
    guidance = {
        "denied": "access denied. Request Reader access and the applicable Cost Management read permissions.",
        "unsupported": "not supported at this scope. Check service availability and billing eligibility.",
        "unavailable": "temporarily unavailable. Retry the report; results may be incomplete.",
    }
    if "coverage" not in result:
        return tuple(clean_text(value) for value in result.get("collectionWarnings", []))
    warnings = []
    for key, item in sorted(result["coverage"].items()):
        status = item.get("status")
        if status in ("ok", "empty"):
            continue
        label = labels.get(key.rsplit("/", 1)[-1], "Data collection")
        subscription = item.get("subscriptionName")
        scope = f"{clean_text(subscription)} - " if subscription else ""
        message = guidance.get(status, "collection incomplete. Retry the report or check read permissions.")
        warnings.append(f"{scope}{label}: {message}")
    return tuple(dict.fromkeys(warnings))


def export_result(result: dict[str, Any]) -> dict[str, Any]:
    exported = redacted_result(result)
    if exported.get("mode") == "connected":
        warnings = collection_warnings(exported)
        exported.pop("coverage", None)
        if warnings:
            exported["collectionWarnings"] = list(warnings)
    return exported


def build_report(result: dict[str, Any]) -> SavingsReport:
    result = redacted_result(result)
    mode = result["mode"]
    connected = mode == "connected"
    records: list[dict[str, Any]] = []
    notices = [clean_text(result.get("notice"))] if result.get("notice") else []
    notices.extend(clean_text(item) for item in result.get("assumptions", []))
    context: list[str] = []
    coverage: list[str] = []
    if connected:
        for key, value in sorted(result.get("context", {}).items()):
            context.append(f"{key}: {clean_text(', '.join(value) if isinstance(value, list) else value)}")
        for key, value in sorted(result.get("filters", {}).items()):
            context.append(f"{key}: {clean_text(value)}")
        coverage.extend(collection_warnings(result))
        for recommendation in result["recommendations"]:
            details = recommendation.get("recommendationDetails", {})
            base = {
                "recordType": "savingsPlan", "label": "Azure Savings Plan recommendation",
                "reference": recommendation.get("recommendationReference"),
                "source": "Azure Cost Management Benefit Recommendations API",
                "subscriptionId": recommendation.get("subscriptionId", result.get("context", {}).get("subscriptionId")),
                "subscriptionName": recommendation.get("subscriptionName"),
                "currency": recommendation.get("currencyCode"), "term": recommendation.get("term"),
                "benefitType": recommendation.get("kind"),
                "benefitSku": recommendation.get("armSkuName"),
                "baselineCost": number(recommendation.get("costWithoutBenefit")),
                "observationHours": number(recommendation.get("totalHours")),
                "notes": f"Granularity: {recommendation.get('commitmentGranularity', 'Hourly')}. "
                         "Amounts apply to Azure's returned observation period, not a full-term forecast.",
            }
            for index, detail in enumerate([details, *recommendation.get("alternatives", [])]):
                records.append({
                    **base,
                    "recordType": "savingsPlan" if index == 0 else "alternative",
                    "label": base["label"] if index == 0 else f"Azure alternative {index}",
                    "hourlyCommitment": number(detail.get("commitmentAmount")),
                    "modeledCost": number(detail.get("totalCost")),
                    "savingsAmount": number(detail.get("savingsAmount")),
                    "savingsPercent": number(detail.get("savingsPercentage")),
                    "coveragePercent": number(detail.get("coveragePercentage")),
                    "utilizationPercent": number(detail.get("averageUtilizationPercentage")),
                    "unusedCommitment": number(detail.get("wastageCost")),
                    "overageCost": number(detail.get("overageCost")),
                    "benefitCost": number(detail.get("benefitCost")),
                })
        for item in result["advisorRecommendations"]:
            resource_id = item.get("resourceId")
            resource = next((
                entry for entry in result.get("computeInventory", [])
                if isinstance(resource_id, str) and entry.get("resourceId", "").lower() == resource_id.lower()
                and entry.get("subscriptionId", result.get("context", {}).get("subscriptionId"))
                == item.get("subscriptionId", result.get("context", {}).get("subscriptionId"))
            ), {})
            records.append({
                "recordType": "advisor", "label": clean_text(item.get("problem")),
                "reference": item.get("recommendationReference"),
                "source": "Azure Advisor Recommendations API", "impact": item.get("impact"),
                "subscriptionId": item.get("subscriptionId", result.get("context", {}).get("subscriptionId")),
                "subscriptionName": item.get("subscriptionName"),
                "computeType": resource.get("computeType"), "computeSku": resource.get("computeSku"),
                "licenseType": resource.get("licenseType"),
                "notes": f"Action: {clean_text(item.get('solution'))}\n"
                         f"Azure last updated: {clean_text(item.get('lastUpdated'))}",
            })
        for entry in result.get("computeInventory", []):
            records.append({
                "recordType": "computeInventory", "label": entry.get("name"),
                "reference": entry.get("recommendationReference"), "source": entry.get("source"),
                "subscriptionId": entry.get("subscriptionId", result.get("context", {}).get("subscriptionId")),
                "subscriptionName": entry.get("subscriptionName"),
                "computeType": entry.get("computeType"), "computeSku": entry.get("computeSku"),
                "licenseType": entry.get("licenseType"),
                "notes": "Current configured inventory, not usage during the recommendation lookback. "
                         "No per-resource discount, Savings Plan eligibility, benefit coverage or savings is inferred. "
                         "A reported license type is configuration evidence only; validate entitlement separately. "
                         "Inventory covers VMs, VM scale sets and App Service plans; other compute services are not collected.",
            })
        for record in records:
            if record["recordType"] in ("savingsPlan", "alternative"):
                inventory = [
                    entry for entry in records if entry["recordType"] == "computeInventory"
                    and entry.get("subscriptionId") == record.get("subscriptionId")
                ]
                record["computeInventory"] = "\n".join(
                    f"{display(entry.get('label'))} | {display(entry.get('computeType'))} | "
                    f"SKU: {display(entry.get('computeSku'))} | License: {display(entry.get('licenseType'))} | "
                    f"Evidence: {display(entry.get('reference'))}"
                    for entry in inventory
                ) or "No inventory records returned. Check source coverage; this does not establish absence of compute."
        notices.extend([
            "Advisor Cost actions are not all Savings Plans. Do not add Advisor and benefit savings together.",
            "Alternatives are mutually exclusive choices, not additive recommendations.",
            "No commitment or savings is inferred when Azure omits a value. Check partial coverage before decisions.",
            "Azure recommendations are cached and may lag workload changes. Review in Azure before purchase.",
            "Compute inventory is current configuration, not attributed benefit coverage, historical usage or verified savings. "
            "VMs, VM scale sets and App Service plans are collected read-only when permitted. "
            "A subscription-level plan cannot be attributed to a resource SKU or per-SKU discount. "
            "Advisor compute details are matched only by exact resource ID.",
        ])
    elif mode == "modeler":
        context.extend([
            f"Currency: {result['currencyCode']}", f"Term: {result['term']} ({result['termHours']} hours)",
            f"Modeled period: {result['periodHours']} elapsed hours; runtime: {result['computeHours']} hours",
            f"PAYG rate: {result['paygRate']} per running hour; discount: {result['discountPercent']}%",
            f"Break-even attainable: {result['breakEvenAttainable']}; positive savings possible: {result['positiveSavingsPossible']}",
        ])
        coverage.append("Offline compute-hours model: user-supplied assumptions, no Azure pricing or live recommendation validation.")
        notices.extend(result["guidance"])
        numeric_fields = (
            "hourlyCommitment", "baselineCost", "modeledCost", "savingsAmount", "savingsPercent",
            "coveragePercent", "utilizationPercent", "unusedCommitment", "remainingPaygCost",
            "projectedTermSavings", "computeHours", "periodHours", "paygRate", "discountPercent",
            "commitmentCost", "breakEvenHours", "breakEvenRuntimePercent", "runtimeMarginHours",
            "projectedTermBaselineCost", "projectedTermModeledCost",
        )
        records.append({
            "recordType": "model", "label": "Compute-hours estimate (selected runtime)",
            "source": result["source"], "currency": result["currencyCode"], "term": result["term"],
            **{key: number(result.get(key)) for key in numeric_fields},
            "observationHours": result["periodHours"],
            "notes": "\n".join([f"Outcome: {result['outcome']}", *result["guidance"], *result["assumptions"]]),
        })
        for point in result["sensitivity"]:
            records.append({
                "recordType": "sensitivity", "label": f"{point['computeHours']:g} running hours",
                "source": result["source"], "currency": result["currencyCode"], "term": result["term"],
                **point, "hourlyCommitment": result["hourlyCommitment"],
                "paygRate": number(result["paygRate"]), "discountPercent": number(result["discountPercent"]),
                "periodHours": result["periodHours"], "observationHours": result["periodHours"],
                "notes": "Independent runtime scenario, not additive. Same rates, discount and elapsed commitment period.",
            })
        notices.append("Offline planning estimate. Period costs are projections, not observed Azure costs. "
                       "Full-term figures repeat the assumed duty cycle; shorter modeling periods do not shorten the plan obligation.")
    else:
        context.extend([f"Currency: {clean_text(result['currencyCode'])}", f"Observed hours: {result['dataHours']}"])
        coverage.append("Offline model: no Azure data collection or live recommendation validation.")
        for option in result["options"] if mode == "comparison" else [result]:
            records.append({
                "recordType": "model",
                "label": option.get("name", f"{option['term']} / {option.get('scenario', '')}"),
                "source": "Local deterministic planning model", "currency": option["currencyCode"],
                "term": option["term"], "hourlyCommitment": number(option.get("hourlyCommitment")),
                "discountPercent": number(option.get("discountPercent")),
                "baselineCost": number(option.get("scenarioEligibleCostObserved")),
                "modeledCost": number(option["modeledCostObserved"]),
                "savingsAmount": number(option["netSavingsObserved"]),
                "savingsPercent": number(option["savingsPercent"]),
                "coveragePercent": number(option["coveragePercent"]),
                "utilizationPercent": number(option["utilizationPercent"]),
                "observationHours": number(option["dataHours"]),
                "unusedCommitment": number(option.get("unusedCommitmentCostObserved")),
                "overageCost": number(option.get("overageCostObserved")),
                "remainingPaygCost": number(option.get("remainingPaygCostObserved")),
                "projectedTermSavings": number(option.get("projectedTermNetSavings")),
                "notes": "\n".join([
                    f"Risk: {option['risk']} | Confidence: {option['confidence']}",
                    *option["warnings"], *option["assumptions"],
                ]),
            })
        notices.append("Offline planning estimate, not an Azure purchase recommendation. "
                       "Amounts cover the supplied observation period, not the full commitment term.")
    for record in records:
        for field, value in record.items():
            if isinstance(value, str):
                record[field] = clean_text(value)
    currencies = sorted({record["currency"] for record in records if record.get("currency")})
    primary = sum(record["recordType"] in ("savingsPlan", "model") for record in records)
    metrics = (
        ("Plan recommendations" if connected else "Modeled options", str(primary)),
        ("Advisor actions", str(sum(record["recordType"] == "advisor" for record in records))),
        ("Data coverage", clean_text(result.get("completeness", "Offline model"))),
        ("Currencies (not summed)", ", ".join(currencies) or "Not supplied"),
    )
    if mode == "modeler":
        metrics = (
            ("Projected period net savings", f"{display(result['savingsAmount'])} {result['currencyCode']}"),
            ("Break-even running hours", display(result["breakEvenHours"]) if result["breakEvenHours"] is not None else "N/A (zero commitment)"),
            ("Break-even feasible", "Yes" if result["breakEvenAttainable"] else "No"),
            ("Modeled outcome", result["outcome"]),
        )
    return SavingsReport(
        title="Savings plans", mode=mode,
        collected_at=clean_text(result.get("collectedAt") or datetime.now(timezone.utc).isoformat()),
        scope=(
            f"{len(result['context']['subscriptionIds'])} selected subscriptions (independent Single scopes)"
            if connected and "subscriptionIds" in result.get("context", {})
            else "Selected subscription (Single)" if connected else
            "One constant-rate workload; user-supplied compute hours" if mode == "modeler" else "User-supplied hourly observations"
        ),
        context=tuple(context), coverage=tuple(coverage), notices=tuple(dict.fromkeys(notices)),
        records=tuple(records), metrics=metrics,
    )


def record_lines(record: dict[str, Any]) -> list[str]:
    return [f"{FIELD_LABELS[field]}: {display(record.get(field))}" for field in REPORT_FIELDS]


def chart_records(report: SavingsReport) -> list[dict[str, Any]]:
    return [
        record for record in report.records
        if record["recordType"] in ("savingsPlan", "model") and record.get("savingsPercent") is not None
    ]


def report_charts(report: SavingsReport) -> list[tuple[str, str, str, list[dict[str, Any]]]]:
    charts = []
    sensitivity = [row for row in report.records if row["recordType"] == "sensitivity"]
    if sensitivity:
        charts.append(("Net savings / loss at different running hours", "savingsAmount",
                       f" {sensitivity[0]['currency']}", sensitivity))
    percentages = chart_records(report)
    if percentages:
        charts.append(("Savings percentage by recommendation / modeled option", "savingsPercent", "%", percentages))
    currencies = sorted({
        row["currency"] for row in report.records
        if row["recordType"] in ("savingsPlan", "model") and row.get("currency")
        and row.get("hourlyCommitment") is not None
    })
    for currency in currencies:
        rows = [
            row for row in report.records
            if row["recordType"] in ("savingsPlan", "model") and row.get("currency") == currency
            and row.get("hourlyCommitment") is not None
        ]
        charts.append((f"Hourly commitment by recommendation / modeled option ({currency}/hour)",
                       "hourlyCommitment", f" {currency}/h", rows))
    return charts


def chart_label(record: dict[str, Any]) -> str:
    scope = record.get("subscriptionName") or record.get("subscriptionId") or ""
    return clean_text(f"{scope} | {record.get('term', '')} | {record['label']}")


def csv_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return "" if value is None else value
    value = clean_text(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_csv(report: SavingsReport, path: Path) -> None:
    columns = ("mode", "collectedAt", "scope", "context", "sourceCoverage", "limitations", *REPORT_FIELDS)
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=columns)
        writer.writeheader()
        metadata = {
            "mode": report.mode, "collectedAt": report.collected_at, "scope": report.scope,
            "context": "\n".join(report.context), "sourceCoverage": "\n".join(report.coverage),
            "limitations": "\n".join(report.notices),
        }
        for record in report.records or ({"recordType": "coverage", "label": "No recommendation records returned"},):
            writer.writerow({key: csv_cell(value) for key, value in {**metadata, **record}.items()})


def export_pdf(report: SavingsReport, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.graphics.shapes import Drawing, Rect, String

    styles = getSampleStyleSheet()
    styles["Title"].textColor = colors.HexColor(NAVY)
    styles["Heading2"].textColor = colors.HexColor(TEAL)

    def paragraph(text: str, style: str = "BodyText") -> Any:
        return Paragraph(escape(text).replace("\n", "<br/>"), styles[style])

    story: list[Any] = [
        paragraph("Savings plans | Executive report", "Title"),
        paragraph("LIVE AZURE RECOMMENDATIONS" if report.mode == "connected" else "OFFLINE PLANNING ESTIMATE", "Heading2"),
        paragraph(f"Collected: {report.collected_at}\nScope: {report.scope}"),
        Spacer(1, 12),
    ]
    cards = Table([[paragraph(label) for label, _ in report.metrics],
                   [paragraph(value, "Heading2") for _, value in report.metrics]], colWidths=[1.75 * inch] * 4)
    cards.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(PALE)),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(TEAL)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend([cards, Spacer(1, 12)])
    charts = report_charts(report)
    for title, field, unit, chart in charts:
        story.append(paragraph(title, "Heading2"))
        drawing = Drawing(540, 175)
        maximum = max(max(abs(row[field]) for row in chart), .01)
        for index, row in enumerate(chart[:5]):
            y = 150 - index * 28
            drawing.add(String(0, y, chart_label(row)[:44], fontSize=8))
            drawing.add(Rect(225, y - 2, 260, 10, fillColor=colors.HexColor("#e4eaf2"), strokeColor=None))
            drawing.add(Rect(225, y - 2, 260 * abs(row[field]) / maximum, 10,
                             fillColor=colors.HexColor(BLUE if row[field] >= 0 else "#b54343"), strokeColor=None))
            drawing.add(String(490, y, f"{row[field]:.2f}{unit}", fontSize=7))
        story.extend([drawing, paragraph(f"Showing {min(5, len(chart))} of {len(chart)} values; alternatives excluded. Length = magnitude; red = negative. Each chart has its own scale. Full records follow.")])
    if not charts:
        story.append(paragraph("No returned numeric values to chart. Review source coverage."))
    story.append(paragraph("Scope, sources and limitations", "Heading2"))
    for text in (*report.context, *report.coverage, *report.notices):
        story.append(paragraph(text))
    story.append(paragraph("Recommendation details", "Heading2"))
    if not report.records:
        story.append(paragraph("No recommendation records returned. Review source coverage; missing values were not estimated."))
    for record in report.records:
        story.append(paragraph(display(record.get("label")), "Heading2"))
        for line in record_lines(record):
            story.append(paragraph(line))
        story.append(Spacer(1, 10))

    def footer(canvas: Any, document: Any) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(TEAL))
        canvas.drawString(36, 22, "Azure Savings Plan Advisor | Decision support only")
        canvas.drawRightString(576, 22, f"Page {document.page}")

    SimpleDocTemplate(
        str(path), pagesize=(612, 792), leftMargin=36, rightMargin=36,
        topMargin=36, bottomMargin=36, title="Savings plans executive report",
    ).build(story, onFirstPage=footer, onLaterPages=footer)


def export_powerpoint(report: SavingsReport, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def slide(title: str, lines: list[str]) -> Any:
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(PALE[1:])
        panel = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.25), Inches(.2), Inches(12.83), Inches(6.65))
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        panel.line.color.rgb = RGBColor.from_string("d2dce5")
        heading = page.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12.3), Inches(.6)).text_frame
        heading.text = title
        heading.paragraphs[0].font.size = Pt(27)
        heading.paragraphs[0].font.bold = True
        heading.paragraphs[0].font.color.rgb = RGBColor.from_string(TEAL[1:])
        body = page.shapes.add_textbox(Inches(.6), Inches(1.15), Inches(12.1), Inches(5.6)).text_frame
        body.word_wrap = True
        for index, line in enumerate(lines):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(15)
            paragraph.font.color.rgb = RGBColor.from_string(NAVY[1:])
            paragraph.space_after = Pt(3)
        footer = page.shapes.add_textbox(Inches(.6), Inches(7), Inches(12), Inches(.25)).text_frame
        footer.text = f"Azure Savings Plan Advisor | Decision support only | {len(presentation.slides)}"
        footer.paragraphs[0].font.size = Pt(10)
        return page

    def pages(title: str, lines: list[str]) -> None:
        wrapped = [
            part for line in lines for part in (
                textwrap.wrap(line, width=105, replace_whitespace=True, break_long_words=True) or [""]
            )
        ]
        for offset in range(0, len(wrapped), 17):
            slide(title + (" (continued)" if offset else ""), wrapped[offset:offset + 17])

    summary = slide("Savings plans | Executive summary", [
        "LIVE AZURE RECOMMENDATIONS" if report.mode == "connected" else "OFFLINE PLANNING ESTIMATE",
        f"Collected: {report.collected_at}", f"Scope: {report.scope}",
    ])
    for index, (label, value) in enumerate(report.metrics):
        x = .6 + index * 3.05
        card = summary.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3), Inches(2.9), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(PALE[1:])
        card.line.color.rgb = RGBColor.from_string("d2dce5")
        title = summary.shapes.add_textbox(Inches(x + .12), Inches(3.15), Inches(2.6), Inches(.4)).text_frame
        title.text = label
        title.paragraphs[0].font.size = Pt(13)
        metric = summary.shapes.add_textbox(Inches(x + .12), Inches(3.6), Inches(2.6), Inches(.6)).text_frame
        metric.text = value
        metric.paragraphs[0].font.size = Pt(20)
        metric.paragraphs[0].font.color.rgb = RGBColor.from_string(TEAL[1:])
    note = summary.shapes.add_textbox(Inches(.6), Inches(5), Inches(12), Inches(.8)).text_frame
    note.word_wrap = True
    note.text = "Financial amounts and alternatives are not summed across records or currencies. Review source coverage, evidence and assumptions before decisions."
    note.paragraphs[0].font.size = Pt(15)
    for title, field, unit, chart in report_charts(report):
        page = slide(title, [])
        maximum = max(max(abs(row[field]) for row in chart), .01)
        for index, row in enumerate(chart[:5]):
            y = 1.5 + index * .85
            label = page.shapes.add_textbox(Inches(.6), Inches(y), Inches(4), Inches(.6)).text_frame
            label.word_wrap = True
            label.text = chart_label(row)[:65]
            label.paragraphs[0].font.size = Pt(13)
            for length, color in [
                (6.5, "#e4eaf2"),
                (max(.001, 6.5 * abs(row[field]) / maximum), BLUE if row[field] >= 0 else "#b54343"),
            ]:
                bar = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.8), Inches(y + .1), Inches(length), Inches(.3))
                bar.fill.solid()
                bar.fill.fore_color.rgb = RGBColor.from_string(color[1:])
                bar.line.fill.background()
            value = page.shapes.add_textbox(Inches(11.5), Inches(y), Inches(1.2), Inches(.5)).text_frame
            value.text = f"{row[field]:.2f}{unit}"
            value.paragraphs[0].font.size = Pt(13)
        note = page.shapes.add_textbox(Inches(.6), Inches(6.2), Inches(12), Inches(.5)).text_frame
        note.text = f"Showing {min(5, len(chart))} of {len(chart)} values; length = magnitude, red = negative. Independent scales; alternatives excluded."
        note.paragraphs[0].font.size = Pt(12)
    pages("Scope, sources and limitations", [*report.context, *report.coverage, *report.notices])
    for record in report.records:
        pages("Recommendation details", record_lines(record))
    if not report.records:
        pages("Recommendation details", ["No recommendation records returned. Review source coverage; no estimate was substituted."])
    presentation.save(str(path))


def write_report(report: SavingsReport, path: Path, kind: str) -> None:
    exporters = {"pdf": export_pdf, "pptx": export_powerpoint, "csv": export_csv}
    if kind not in exporters:
        raise ValueError("Unsupported report format.")
    descriptor, temporary = tempfile.mkstemp(dir=path.parent, suffix=f".{kind}")
    os.close(descriptor)
    temporary_path = Path(temporary)
    try:
        exporters[kind](report, temporary_path)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)
