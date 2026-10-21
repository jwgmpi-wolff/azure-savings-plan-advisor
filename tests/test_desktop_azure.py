import json
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
import requests

from api import azure_provider, desktop_azure, desktop_identity
from api.azure_provider import ConnectedProviderError, arm_list
from api.desktop_azure import DesktopAzure


TENANT = "00000000-0000-0000-0000-000000000001"
OTHER = "00000000-0000-0000-0000-000000000002"
SUB = "00000000-0000-0000-0000-000000000003"
TOKEN = "synthetic-test-token"


class Credential:
    def __init__(self, **kwargs):
        self.tenant = kwargs["tenant_id"]

    def get_token(self, scope):
        assert scope == azure_provider.MANAGEMENT_RESOURCE
        return SimpleNamespace(token=TOKEN)

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return None


@pytest.fixture
def client(monkeypatch):
    client = DesktopAzure()
    monkeypatch.setattr(client, "current_tenant", lambda: TENANT)
    monkeypatch.setattr(desktop_azure, "SilentAzureCliCredential", Credential)

    def listing(url, token, params, source):
        assert token == TOKEN
        if url.endswith("/tenants"):
            return [{"tenantId": TENANT}, {"tenantId": OTHER}]
        if url.endswith("/subscriptions"):
            return [
                {"tenantId": TENANT, "subscriptionId": SUB, "displayName": "Test subscription", "state": "Enabled"},
                {"tenantId": OTHER, "subscriptionId": OTHER, "state": "Enabled"},
                {"tenantId": TENANT, "subscriptionId": OTHER, "state": "Disabled"},
            ]
        if any(url.endswith("/" + resource_type) for _key, resource_type, _version, _label in desktop_azure.COMPUTE_SOURCES):
            return []
        return [{
            "id": "/synthetic/recommendation", "properties": {
                "category": "Cost", "impact": "High",
                "shortDescription": {"problem": "Savings Plan opportunity", "solution": "Review plan"},
                "extendedProperties": {"private": "must not be forwarded"},
            },
        }]
    monkeypatch.setattr(desktop_azure, "arm_list", listing)
    monkeypatch.setattr(desktop_azure, "get_benefit_recommendations", lambda *args, **kwargs: {
        "recommendations": [{
            "term": "P1Y", "currencyCode": "USD", "commitmentGranularity": "Hourly",
            "recommendationDetails": {"commitmentAmount": 4.25, "savingsPercentage": 20},
        }],
        "recommendationCount": 1, "source": "Azure Cost Management",
        "apiVersion": azure_provider.API_VERSION, "notice": "Azure returned recommendations.",
    })
    return client


def test_context_lists_only_enabled_subscriptions_in_selected_tenant(client):
    context = client.refresh(TENANT)
    assert context["tenantId"] == TENANT
    assert context["tenants"] == [TENANT, OTHER]
    assert [item["subscriptionId"] for item in context["subscriptions"]] == [SUB]


def test_connected_analysis_needs_no_csv_and_preserves_azure_commitment(client):
    result = client.analyze(TENANT, SUB, "P1Y", "Last30Days")
    assert result["mode"] == "connected"
    assert result["completeness"] == "complete"
    assert result["recommendations"][0]["recommendationDetails"]["commitmentAmount"] == 4.25
    assert result["coverage"]["advisor"]["status"] == "ok"
    assert result["context"] == {"tenantId": TENANT, "subscriptionId": SUB}
    assert "extendedProperties" not in json.dumps(result)
    assert "collectedAt" in result


def test_cross_tenant_or_disabled_subscription_is_rejected_before_collection(client, monkeypatch):
    benefits = Mock()
    monkeypatch.setattr(desktop_azure, "get_benefit_recommendations", benefits)
    with pytest.raises(ValueError, match="not enabled or accessible"):
        client.analyze(TENANT, OTHER, "P1Y", "Last30Days")
    benefits.assert_not_called()


def test_unknown_tenant_is_rejected_and_previous_context_cleared(client):
    client.refresh(TENANT)
    with pytest.raises(ConnectedProviderError, match="not available"):
        client.refresh(SUB)
    assert client.tenant_id is None
    assert client.subscriptions == []


def test_permission_failure_is_partial_not_empty_or_offline_fallback(client, monkeypatch):
    def denied(*args, **kwargs):
        raise ConnectedProviderError("Cost Management denied. Request Cost Management Reader.", status="denied")
    monkeypatch.setattr(desktop_azure, "get_benefit_recommendations", denied)
    result = client.analyze(TENANT, SUB, "P1Y", "Last30Days")
    assert result["completeness"] == "partial"
    assert result["coverage"]["savingsPlan"]["status"] == "denied"
    assert result["recommendations"] == []
    assert result["advisorRecommendations"]


def test_empty_sources_stay_empty(client, monkeypatch):
    original = desktop_azure.arm_list
    monkeypatch.setattr(desktop_azure, "arm_list", lambda url, *args: [] if "/Microsoft.Advisor/" in url else original(url, *args))
    monkeypatch.setattr(desktop_azure, "get_benefit_recommendations", lambda *args, **kwargs: {
        "recommendations": [], "recommendationCount": 0, "source": "Azure",
        "apiVersion": azure_provider.API_VERSION, "notice": "No recommendations.",
    })
    result = client.analyze(TENANT, SUB, "P3Y", "Last60Days")
    assert result["completeness"] == "complete"
    assert {item["status"] for item in result["coverage"].values()} == {"empty"}


def test_multiple_subscriptions_keep_results_and_coverage_separate(client, monkeypatch, tmp_path):
    from api.reporting import build_report, write_report
    from pypdf import PdfReader
    from pptx import Presentation

    def refresh(tenant):
        client.subscriptions = [
            {"subscriptionId": SUB, "name": "Synthetic Finance"},
            {"subscriptionId": OTHER, "name": "Synthetic Engineering"},
        ]
    monkeypatch.setattr(client, "refresh", refresh)
    original = desktop_azure.get_benefit_recommendations

    def benefits(scope, *args, **kwargs):
        if OTHER in scope:
            raise ConnectedProviderError("Cost access denied.", status="denied")
        return original(scope, *args, **kwargs)
    monkeypatch.setattr(desktop_azure, "get_benefit_recommendations", benefits)
    result = client.analyze_many(TENANT, [SUB, OTHER, SUB], "P1Y", "Last30Days")
    assert result["context"]["subscriptionIds"] == sorted([SUB, OTHER])
    assert result["completeness"] == "partial"
    assert len(result["coverage"]) == 10
    assert result["coverage"][f"{OTHER}/savingsPlan"]["status"] == "denied"
    assert result["recommendations"][0]["subscriptionId"] == SUB
    assert len(result["advisorRecommendations"]) == 2
    report = build_report(result)
    assert "2 selected subscriptions" in report.scope
    assert {item["subscriptionId"] for item in report.records} == {SUB, OTHER}
    assert all(item["subscriptionName"] for item in report.records)
    for kind in ("pdf", "pptx", "csv"):
        path = tmp_path / f"multiple.{kind}"
        write_report(report, path, kind)
        if kind == "pdf":
            text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
        elif kind == "pptx":
            text = "\n".join(
                shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame
            )
        else:
            text = path.read_text(encoding="utf-8-sig")
        assert SUB in text and OTHER in text
        assert "Synthetic Finance" in text and "Synthetic Engineering" in text
        assert "denied" in text


@pytest.mark.parametrize("selected", [[], [SUB, OTHER]])
def test_invalid_batch_selection_collects_no_recommendations(client, monkeypatch, selected):
    collect = Mock()
    monkeypatch.setattr(client, "analyze", collect)
    with pytest.raises(ValueError):
        client.analyze_many(TENANT, selected, "P1Y", "Last30Days")
    collect.assert_not_called()

def test_compute_collection_uses_delegated_scope_and_safe_fields(client, monkeypatch):
    original = desktop_azure.arm_list
    calls = []

    def listing(url, token, params, source):
        for key, resource_type, version, _label in desktop_azure.COMPUTE_SOURCES:
            if url.endswith("/" + resource_type):
                calls.append((url, token, params))
                assert url.startswith(f"https://management.azure.com/subscriptions/{SUB}/providers/")
                assert params == {"api-version": version}
                assert token == TOKEN
                return [{
                    "id": f"/subscriptions/{SUB}/resourceGroups/synthetic/providers/{resource_type}/synthetic-{key}",
                    "type": resource_type, "name": f"synthetic-{key}",
                    "sku": {"name": "P1v3" if key == "appServicePlans" else "Standard_D4s_v5"},
                    "properties": {
                        "hardwareProfile": {"vmSize": "Standard_D2s_v5"},
                        "licenseType": "Windows_Server",
                        "virtualMachineProfile": {"licenseType": "RHEL_BYOS"},
                        "osProfile": {"adminPassword": "synthetic-private-password"},
                        "customData": "must not be forwarded",
                    },
                }]
        return original(url, token, params, source)
    monkeypatch.setattr(desktop_azure, "arm_list", listing)
    result = client.analyze_many(TENANT, [SUB], "P1Y", "Last30Days")
    assert len(calls) == 3
    inventory = result["computeInventory"]
    assert len(inventory) == 3
    assert [entry["computeSku"] for entry in inventory] == ["Standard_D2s_v5", "Standard_D4s_v5", "P1v3"]
    assert inventory[1]["licenseType"] == "RHEL_BYOS"
    assert all(entry["subscriptionId"] == SUB for entry in inventory)
    assert all(entry["subscriptionName"] == "Test subscription" for entry in inventory)
    assert all(entry["recommendationReference"] for entry in inventory)
    text = json.dumps(result)
    assert "synthetic-private-password" not in text
    assert "customData" not in text
    assert result["completeness"] == "complete"


@pytest.mark.parametrize("status", ["denied", "unavailable", "unsupported"])
def test_compute_failure_keeps_recommendations_and_other_inventory(client, monkeypatch, status):
    original = desktop_azure.arm_list

    def listing(url, *args):
        if url.endswith("/Microsoft.Compute/virtualMachines"):
            raise ConnectedProviderError("Inventory read failed. Request Reader access.", status=status)
        return original(url, *args)
    monkeypatch.setattr(desktop_azure, "arm_list", listing)
    result = client.analyze(TENANT, SUB, "P1Y", "Last30Days")
    assert result["completeness"] == "partial"
    assert result["coverage"]["virtualMachines"]["status"] == status
    assert result["coverage"]["appServicePlans"]["status"] == "empty"
    assert result["recommendations"]
    assert result["computeInventory"] == []


@pytest.mark.parametrize("change", [
    {"id": f"/subscriptions/{OTHER}/resourceGroups/test/providers/Microsoft.Compute/virtualMachines/vm"},
    {"type": "Microsoft.Web/sites"},
    {"properties": None},
])
def test_invalid_inventory_is_partial_not_clean(client, monkeypatch, change):
    original = desktop_azure.arm_list

    def listing(url, *args):
        if url.endswith("/Microsoft.Compute/virtualMachines"):
            return [{
                "id": f"/subscriptions/{SUB}/resourceGroups/test/providers/Microsoft.Compute/virtualMachines/vm",
                "type": "Microsoft.Compute/virtualMachines", "properties": {}, **change,
            }]
        return original(url, *args)
    monkeypatch.setattr(desktop_azure, "arm_list", listing)
    result = client.analyze(TENANT, SUB, "P1Y", "Last30Days")
    assert result["completeness"] == "partial"
    assert result["coverage"]["virtualMachines"]["status"] == "unavailable"
    assert "malformed or out-of-scope" in result["coverage"]["virtualMachines"]["message"]
    assert result["computeInventory"] == []


def test_one_subscription_failure_does_not_discard_successful_results(client, monkeypatch):
    def refresh(tenant):
        client.subscriptions = [{"subscriptionId": value, "name": "Synthetic"} for value in [SUB, OTHER]]
    monkeypatch.setattr(client, "refresh", refresh)
    original = client.analyze

    def collect(tenant, subscription, *args):
        if subscription == OTHER:
            raise ConnectedProviderError("Authentication unavailable.")
        return original(tenant, subscription, *args)
    monkeypatch.setattr(client, "analyze", collect)
    result = client.analyze_many(TENANT, [SUB, OTHER], "P1Y", "Last30Days")
    assert result["completeness"] == "partial"
    assert result["recommendations"][0]["subscriptionId"] == SUB
    assert result["coverage"][f"{OTHER}/advisor"]["status"] == "unavailable"
    assert result["coverage"][f"{OTHER}/savingsPlan"]["status"] == "unavailable"


def test_cli_sign_in_is_tenant_specific_and_does_not_expose_output(monkeypatch):
    client = DesktopAzure()
    monkeypatch.setattr(desktop_identity.shutil, "which", lambda _: "az.cmd")
    calls = []

    def run(args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(returncode=0, stdout=json.dumps({"tenantId": TENANT, "userType": "user"}).encode())
    monkeypatch.setattr(desktop_identity.subprocess, "run", run)
    assert client.sign_in(TENANT) == TENANT
    assert calls[0][0][1:4] == ["login", "--tenant", TENANT]
    assert calls[0][1]["env"]["AZURE_CORE_LOGIN_EXPERIENCE_V2"] == "off"
    assert calls[0][1]["capture_output"] is True
    assert calls[0][1]["check"] is False


def test_cli_service_principal_is_never_used(monkeypatch):
    client = DesktopAzure()
    monkeypatch.setattr(desktop_identity.shutil, "which", lambda _: "az.cmd")
    monkeypatch.setattr(desktop_identity.subprocess, "run", lambda *args, **kwargs: SimpleNamespace(
        returncode=0, stdout=json.dumps({"tenantId": TENANT, "userType": "servicePrincipal"}).encode(),
    ))
    with pytest.raises(ConnectedProviderError, match="interactive user"):
        client.refresh(TENANT)
    assert client.subscriptions == []


@pytest.mark.parametrize("tenant", ["invalid", TENANT + " & whoami"])
def test_tenant_input_cannot_inject_cli_arguments(monkeypatch, tenant):
    run = Mock()
    monkeypatch.setattr(desktop_identity.subprocess, "run", run)
    with pytest.raises(ValueError, match="GUID"):
        DesktopAzure().sign_in(tenant)
    run.assert_not_called()


def test_arm_pagination_keeps_credentials_on_selected_scope(monkeypatch):
    url = "https://management.azure.com/subscriptions"
    calls = []

    def get(target, **kwargs):
        calls.append((target, kwargs))
        payload = {"value": [{"page": len(calls)}]}
        if len(calls) == 1:
            payload["nextLink"] = url + "?continuation=2"
        return SimpleNamespace(status_code=200, raise_for_status=lambda: None, json=lambda: payload)
    monkeypatch.setattr(azure_provider.requests, "get", get)
    assert arm_list(url, TOKEN, {"api-version": "2022-12-01"}, "Test") == [{"page": 1}, {"page": 2}]
    assert calls[1][1]["params"] == {}
    assert calls[0][1]["allow_redirects"] is False


@pytest.mark.parametrize("next_link", [
    "https://attacker.invalid/subscriptions", "http://management.azure.com/subscriptions",
    "https://management.azure.com/subscriptions/other", "https://management.azure.com/subscriptions",
])
def test_untrusted_or_repeated_continuation_never_receives_token(monkeypatch, next_link):
    get = Mock(return_value=SimpleNamespace(
        status_code=200, raise_for_status=lambda: None,
        json=lambda: {"value": [], "nextLink": next_link},
    ))
    monkeypatch.setattr(azure_provider.requests, "get", get)
    with pytest.raises(ConnectedProviderError, match="continuation"):
        arm_list("https://management.azure.com/subscriptions", TOKEN, {}, "Test")
    assert get.call_count == 1


@pytest.mark.parametrize("status,state", [(401, "denied"), (403, "denied"), (404, "unsupported"), (429, "unavailable")])
def test_arm_errors_have_distinct_redacted_states(monkeypatch, status, state):
    monkeypatch.setattr(azure_provider.requests, "get", lambda *args, **kwargs: SimpleNamespace(status_code=status))
    with pytest.raises(ConnectedProviderError) as failure:
        arm_list("https://management.azure.com/subscriptions", TOKEN, {}, "Test")
    assert failure.value.status == state
    assert TOKEN not in str(failure.value)


def test_network_failure_does_not_echo_sensitive_exception(monkeypatch):
    def fail(*args, **kwargs):
        raise requests.ConnectionError("secret diagnostic " + TOKEN)
    monkeypatch.setattr(azure_provider.requests, "get", fail)
    with pytest.raises(ConnectedProviderError) as failure:
        arm_list("https://management.azure.com/subscriptions", TOKEN, {}, "Test")
    assert TOKEN not in str(failure.value)
