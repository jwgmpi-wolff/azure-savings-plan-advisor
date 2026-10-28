from unittest.mock import Mock, patch

import pytest
import requests

from api.compute_modeler import calculate_compute_model
from api.copilot_provider import AI_SOURCE, DEFAULT_MODEL, CopilotError, explain_model
from api.reporting import build_report, export_result, write_report


def model(**overrides):
    return calculate_compute_model(**{
        "period_hours": 100, "compute_hours": 90, "payg_rate": 10, "discount_percent": 20,
        "hourly_commitment": 8, "currency_code": "USD", "term": "P1Y", **overrides,
    })


@pytest.fixture
def service():
    with patch("api.copilot_provider.requests.Session") as factory:
        client = factory.return_value.__enter__.return_value
        metadata = Mock(status_code=200)
        metadata.json.return_value = {"model_info": {"general.architecture": "llama"}, "capabilities": ["completion"]}
        reply = Mock(status_code=200)
        reply.json.return_value = {
            "model": DEFAULT_MODEL, "done": True, "done_reason": "stop",
            "message": {"role": "assistant", "content": "Break-even is 80 running hours (80%). Savings are 100 USD."},
        }
        client.post.side_effect = [metadata, reply]
        yield factory, client, metadata, reply


def test_local_request_contract_minimizes_data_and_never_uses_proxy_or_credentials(service):
    factory, client, _, _ = service
    result = model()
    result["context"] = {"tenantId": "synthetic-private-tenant"}
    result["copilotAnalysis"] = {"explanation": "synthetic-prior-response"}
    with patch.dict("os.environ", {"HTTP_PROXY": "http://synthetic.invalid", "GITHUB_TOKEN": "synthetic-secret"}):
        analysis = explain_model(result)
    assert analysis["status"] == "complete"
    assert analysis["source"] == AI_SOURCE
    assert analysis["model"] == DEFAULT_MODEL
    assert analysis["modelCollectedAt"] == result["collectedAt"]
    assert "80 running hours" in analysis["explanation"]
    assert client.trust_env is False
    show, chat = client.post.call_args_list
    assert show.args == ("http://127.0.0.1:11434/api/show",)
    assert show.kwargs["json"] == {"model": DEFAULT_MODEL}
    assert chat.args == ("http://127.0.0.1:11434/api/chat",)
    assert chat.kwargs["timeout"] == (3, 150)
    assert all(call.kwargs["allow_redirects"] is False for call in (show, chat))
    payload = chat.kwargs["json"]
    assert payload["stream"] is False and payload["model"] == DEFAULT_MODEL
    assert "tools" not in payload
    assert payload["messages"][0]["role"] == "system"
    prompt = payload["messages"][1]["content"]
    assert '"breakEvenHours": 80.0' in prompt
    for excluded in ("synthetic-private-tenant", "synthetic-prior-response", "synthetic-secret"):
        assert excluded not in str(payload)
    factory.return_value.__exit__.assert_called_once()


@pytest.mark.parametrize("name", ["", " ", "../bad model", "access_token=secret", "x\ninjected"])
def test_invalid_model_blocks_network(service, name):
    factory, _, _, _ = service
    with pytest.raises(CopilotError, match="installed local"):
        explain_model(model(), model=name)
    factory.assert_not_called()


@pytest.mark.parametrize("name", ["synthetic:cloud", "synthetic:Cloud-long"])
def test_cloud_model_names_block_network(service, name):
    factory, _, _, _ = service
    with pytest.raises(CopilotError, match="Cloud models"):
        explain_model(model(), model=name)
    factory.assert_not_called()


@pytest.mark.parametrize("metadata", [
    {"remote_host": "https://synthetic.invalid", "model_info": {"local": True}},
    {"remote_model": "synthetic-remote", "model_info": {"local": True}},
    {"model_info": {}}, {"model_info": True}, {"model_info": ["invalid"]}, {},
])
def test_remote_or_unverifiable_model_blocks_financial_transmission(service, metadata):
    _, client, show, _ = service
    show.json.return_value = metadata
    with pytest.raises(CopilotError, match="verified as local"):
        explain_model(model())
    assert client.post.call_count == 1
    assert "messages" not in client.post.call_args.kwargs["json"]


def test_embedding_only_model_blocks_financial_transmission(service):
    _, client, show, _ = service
    show.json.return_value = {"model_info": {"local": True}, "capabilities": ["embedding"]}
    with pytest.raises(CopilotError, match="does not support chat"):
        explain_model(model())
    assert client.post.call_count == 1


def test_missing_model_has_actionable_error_without_auto_download(service):
    _, client, show, _ = service
    show.status_code = 404
    with pytest.raises(CopilotError, match="ollama list"):
        explain_model(model())
    assert client.post.call_count == 1


@pytest.mark.parametrize("status", [301, 302, 307, 401, 403, 429, 500])
def test_http_errors_and_redirects_are_not_successes(service, status):
    _, client, show, _ = service
    show.status_code = status
    with pytest.raises(CopilotError, match=f"HTTP {status}"):
        explain_model(model())
    assert client.post.call_count == 1


@pytest.mark.parametrize("failure,expected", [
    (requests.Timeout("access_token=private-secret"), "timed out"),
    (requests.ConnectionError("access_token=private-secret"), "Start Ollama"),
    (requests.RequestException("access_token=private-secret"), "request failed"),
])
def test_connection_failures_are_safe_not_success_shaped(service, failure, expected, caplog):
    _, client, _, _ = service
    client.post.side_effect = failure
    with pytest.raises(CopilotError, match=expected) as error:
        explain_model(model())
    assert "private-secret" not in str(error.value)
    assert "private-secret" not in caplog.text


@pytest.mark.parametrize("stage", [2, 1])
def test_invalid_json_is_reported_at_either_stage(service, stage):
    _, _, show, reply = service
    (show if stage == 1 else reply).json.side_effect = ValueError("private-secret")
    with pytest.raises(CopilotError, match="invalid JSON") as error:
        explain_model(model())
    assert "private-secret" not in str(error.value)


@pytest.mark.parametrize("response", [[], None, {"error": "access_token=private-secret"}])
def test_failed_response_shapes_are_not_successes(service, response):
    _, _, _, reply = service
    reply.json.return_value = response
    with pytest.raises(CopilotError, match="invalid or failed") as error:
        explain_model(model())
    assert "private-secret" not in str(error.value)


@pytest.mark.parametrize("changes", [
    {"done": False}, {"done_reason": "length"}, {"model": "unexpected-model"},
    {"message": None}, {"message": {"role": "user", "content": "untrusted"}},
    *({"message": {"role": "assistant", "content": content}} for content in [None, "", " ", 42, "x" * 20001]),
    {"message": {"role": "assistant", "content": "tool request", "tool_calls": [{"function": "synthetic"}]}},
])
def test_empty_truncated_tool_or_wrong_model_responses_are_rejected(service, changes):
    _, _, _, reply = service
    reply.json.return_value.update(changes)
    with pytest.raises(CopilotError, match="no complete, verifiable explanation"):
        explain_model(model())


@pytest.mark.parametrize("overrides,figure", [
    ({"hourly_commitment": 0}, '"breakEvenHours": null'),
    ({"hourly_commitment": 12}, '"breakEvenAttainable": false'),
    ({"compute_hours": 50}, '"savingsAmount": -300.0'),
])
def test_edge_cases_are_sent_without_fabrication(service, overrides, figure):
    _, client, _, _ = service
    explain_model(model(**overrides))
    assert figure in client.post.call_args.kwargs["json"]["messages"][1]["content"]


def test_ai_response_is_redacted_and_preserved_with_provenance_in_all_exports(service, tmp_path):
    _, _, _, reply = service
    reply.json.return_value["message"]["content"] = "Synthetic explanation. access_token=private-secret"
    result = model()
    result["copilotAnalysis"] = explain_model(result)
    report = build_report(result)
    assert "private-secret" not in str(export_result(result))
    assert any(AI_SOURCE in notice and DEFAULT_MODEL in notice for notice in report.notices)
    for kind in ("csv", "pdf", "pptx"):
        path = tmp_path / f"report.{kind}"
        write_report(report, path, kind)
        if kind == "csv":
            content = path.read_text(encoding="utf-8-sig")
        elif kind == "pdf":
            from pypdf import PdfReader
            content = "\n".join(page.extract_text() for page in PdfReader(path).pages)
        else:
            from pptx import Presentation
            content = "\n".join(
                shape.text for slide in Presentation(path).slides
                for shape in slide.shapes if shape.has_text_frame
            )
        assert "Synthetic explanation." in content
        assert AI_SOURCE in content and DEFAULT_MODEL in content
        assert "private-secret" not in content
