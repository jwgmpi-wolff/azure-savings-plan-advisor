import copy
import csv
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from api.reporting import (
    build_report, chart_records, clean_text, export_csv, export_pdf,
    collection_warnings, export_powerpoint, export_result, redacted_result, write_report, report_charts,
)
from desktop_app import analyze_csv_file
from api.compute_modeler import calculate_compute_model


@pytest.fixture
def connected_result():
    return {
        "mode": "connected", "collectedAt": "2026-01-01T00:00:00Z",
        "completeness": "partial", "context": {"tenantId": "synthetic-tenant", "subscriptionId": "synthetic-subscription"},
        "filters": {"term": "P1Y", "lookbackPeriod": "Last30Days", "recommendationScope": "Single"},
        "coverage": {
            "advisor": {"status": "denied", "source": "Azure Advisor", "apiVersion": "2025-01-01", "message": "Request Reader access."},
            "savingsPlan": {"status": "ok", "source": "Azure Cost Management", "apiVersion": "2026-06-01"},
        },
        "notice": "Synthetic fixture only. Alternatives are not additive.",
        "recommendations": [{
            "recommendationReference": "evidence-001", "term": "P1Y",
            "currencyCode": "USD", "commitmentGranularity": "Hourly",
            "costWithoutBenefit": 1000, "totalHours": 720,
            "recommendationDetails": {
                "commitmentAmount": 4.25, "savingsAmount": 123.45,
                "savingsPercentage": 12.345, "coveragePercentage": 80,
                "averageUtilizationPercentage": 95, "totalCost": 876.55,
            },
            "alternatives": [{"commitmentAmount": 2.75, "savingsAmount": 100, "savingsPercentage": 10}],
        }],
        "advisorRecommendations": [],
    }


def text_from_pdf(path):
    return "\n".join(page.extract_text() for page in PdfReader(path).pages)


def text_from_powerpoint(path):
    return "\n".join(
        shape.text for slide in Presentation(path).slides
        for shape in slide.shapes if shape.has_text_frame
    )


@pytest.mark.parametrize("kind,reader", [
    ("pdf", text_from_pdf), ("pptx", text_from_powerpoint),
    ("csv", lambda path: path.read_text(encoding="utf-8-sig")),
])
def test_compute_model_exports_preserve_break_even_guidance_and_sensitivity(tmp_path, kind, reader):
    result = calculate_compute_model(
        period_hours=100, compute_hours=90, payg_rate=10, discount_percent=20,
        hourly_commitment=8, currency_code="USD", term="P1Y",
    )
    report = build_report(result)
    assert report.mode == "modeler"
    assert report.metrics[0] == ("Projected period net savings", "100.00 USD")
    assert report.records[0]["breakEvenHours"] == 80
    assert report.records[0]["projectedTermSavings"] == 8760
    assert len(report.records) == 6
    sensitivity = [row for row in report.records if row["recordType"] == "sensitivity"]
    assert [row["computeHours"] for row in sensitivity] == [0, 50, 80, 90, 100]
    charts = report_charts(report)
    assert charts[0][1] == "savingsAmount"
    assert charts[0][3][0]["savingsAmount"] == -800
    assert len(charts) == 3
    path = tmp_path / f"compute.{kind}"
    write_report(report, path, kind)
    text = reader(path)
    for expected in ("Calculation-based reference figures", "Break-even", "80", "8760" if kind == "csv" else "8,760",
                     "100 running hours", "duty cycle", "idle"):
        assert expected in text
    if kind != "csv":
        assert "Net savings / loss" in text


def test_report_preserves_azure_values_and_missing_is_not_zero(connected_result):
    report = build_report(connected_result)
    assert report.collected_at == connected_result["collectedAt"]
    assert report.records[0]["hourlyCommitment"] == 4.25
    assert report.records[0]["savingsAmount"] == 123.45
    assert report.records[1]["recordType"] == "alternative"
    assert report.records[1]["modeledCost"] is None
    assert report.metrics[0][1] == "1"
    assert report.metrics[2][1] == "partial"
    assert len(chart_records(report)) == 1
    assert any("incomplete" in line for line in report.coverage)

@pytest.mark.parametrize("kind,reader", [
    ("pdf", text_from_pdf), ("pptx", text_from_powerpoint),
    ("csv", lambda path: path.read_text(encoding="utf-8-sig")),
])
@pytest.mark.parametrize("partial", [False, True])
def test_reports_suppress_raw_collection_diagnostics_but_keep_failures(connected_result, tmp_path, kind, reader, partial):
    connected_result["coverage"] = {
        "synthetic-subscription/advisor": {
            "status": "denied" if partial else "ok", "source": "technical-source-marker",
            "apiVersion": "2025-01-01", "message": "verbose-diagnostic-marker",
            "subscriptionName": "Synthetic Finance",
        },
        "synthetic-subscription/savingsPlan": {"status": "ok", "apiVersion": "2026-06-01"},
        "synthetic-subscription/virtualMachines": {"status": "empty", "apiVersion": "2024-07-01"},
    }
    connected_result["completeness"] = "partial" if partial else "complete"
    if partial:
        for index in range(70):
            connected_result["coverage"][f"synthetic-subscription-{index}/savingsPlan"] = {
                "status": "denied", "subscriptionName": f"Synthetic failed subscription {index}",
            }
    report = build_report(connected_result)
    assert len(report.coverage) == (1 if partial else 0)
    path = tmp_path / f"quiet.{kind}"
    write_report(report, path, kind)
    text = reader(path)
    for unwanted in (
        "technical-source-marker", "verbose-diagnostic-marker", "2025-01-01",
        "2026-06-01", "2024-07-01", "synthetic-subscription/advisor",
        "synthetic-subscription/virtualMachines", "synthetic-subscription/savingsPlan",
    ):
        assert unwanted not in text
    assert "evidence-001" in text
    assert "123.45" in text
    if partial:
        assert "incomplete" in text
        assert "Synthetic Finance" not in text
        assert "Synthetic failed subscription" not in text
        assert "access denied" not in text and "Request Reader access" not in text
    else:
        assert "access denied" not in text
    if kind == "csv":
        rows = list(csv.DictReader(path.open(encoding="utf-8-sig", newline="")))
        assert bool(rows[0]["sourceCoverage"]) == partial


@pytest.mark.parametrize("status,expected", [
    ("ok", None), ("empty", None), ("denied", "access denied"),
    ("unsupported", "not supported"), ("unavailable", "temporarily unavailable"),
    ("unexpected", "collection incomplete"),
])
def test_json_export_omits_log_entries_and_preserves_incomplete_summary(connected_result, status, expected):
    connected_result["coverage"] = {
        "synthetic-subscription/virtualMachines": {
            "status": status, "apiVersion": "2024-07-01", "message": "raw-diagnostic",
        },
    }
    connected_result["completeness"] = "complete" if status in ("ok", "empty") else "partial"
    original = copy.deepcopy(connected_result)
    exported = export_result(connected_result)
    assert "coverage" not in exported
    assert connected_result == original
    assert exported["recommendations"] == connected_result["recommendations"]
    assert "collectionWarnings" not in exported
    if expected:
        assert "incomplete" in exported["collectionSummary"]
        assert expected in collection_warnings(connected_result)[0]
    else:
        assert "collectionSummary" not in exported
    assert build_report(exported).coverage == build_report(connected_result).coverage

@pytest.mark.parametrize("legacy_entries", [[], ["Synthetic Finance - access denied. Request Reader access."]])
def test_partial_json_keeps_summary_without_reexporting_legacy_log_entries(connected_result, legacy_entries):
    connected_result.pop("coverage")
    connected_result["collectionWarnings"] = legacy_entries
    connected_result["completeness"] = "partial"
    exported = export_result(connected_result)
    assert "collectionWarnings" not in exported and "coverage" not in exported
    assert exported["collectionSummary"] == "Data collection is incomplete. Some data could not be collected."
    assert build_report(exported).coverage == (exported["collectionSummary"],)
    assert "Synthetic Finance" not in str(exported)

def test_report_preserves_benefit_metadata_without_inventing_compute_or_discount(connected_result, tmp_path):
    recommendation = connected_result["recommendations"][0]
    recommendation.update(kind="SavingsPlan", armSkuName="Synthetic_Plan_SKU")
    recommendation["recommendationDetails"]["benefitCost"] = 700
    report = build_report(connected_result)
    primary, alternative = report.records
    assert primary["benefitType"] == alternative["benefitType"] == "SavingsPlan"
    assert primary["benefitSku"] == alternative["benefitSku"] == "Synthetic_Plan_SKU"
    assert primary["benefitCost"] == 700
    assert alternative["benefitCost"] is None
    assert primary.get("computeType") is None
    assert primary.get("computeSku") is None
    assert primary.get("discountPercent") is None
    path = tmp_path / "benefits.csv"
    export_csv(report, path)
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig", newline="")))
    assert rows[0]["benefitSku"] == "Synthetic_Plan_SKU"
    assert rows[0]["benefitCost"] == "700.0"
    assert rows[0]["computeSku"] == rows[0]["discountPercent"] == ""

@pytest.mark.parametrize("kind,reader", [
    ("pdf", text_from_pdf), ("pptx", text_from_powerpoint),
    ("csv", lambda path: path.read_text(encoding="utf-8-sig")),
])
def test_inventory_is_scoped_and_exported_without_inventing_benefits(connected_result, tmp_path, kind, reader):
    resource_id = "/subscriptions/synthetic-subscription/resourceGroups/test/providers/Microsoft.Compute/virtualMachines/vm"
    connected_result["computeInventory"] = [
        {
            "resourceId": resource_id, "subscriptionId": "synthetic-subscription",
            "name": "Synthetic workload", "computeType": "Virtual machine", "computeSku": "Standard_D2s_v5",
            "licenseType": "Windows_Server", "recommendationReference": "inventory-001",
            "source": "Azure Resource Manager Microsoft.Compute/virtualMachines List",
        },
        {
            "resourceId": "/subscriptions/other/resourceGroups/test/providers/Microsoft.Web/serverfarms/plan",
            "subscriptionId": "other", "name": "Other subscription workload",
            "computeType": "App Service plan", "computeSku": "P1v3",
        },
    ]
    connected_result["advisorRecommendations"] = [
        {"resourceId": resource_id.upper(), "problem": "Review synthetic workload"},
        {"resourceId": resource_id + "-other", "problem": "Unmatched action"},
    ]
    report = build_report(connected_result)
    plan, alternative, matched, unmatched, vm, app = report.records
    assert vm["computeSku"] == "Standard_D2s_v5"
    assert vm["licenseType"] == "Windows_Server"
    assert vm.get("savingsAmount") is None
    assert vm.get("discountPercent") is None
    assert plan.get("computeSku") is None
    assert "Standard_D2s_v5" in plan["computeInventory"]
    assert plan["computeInventory"] == alternative["computeInventory"]
    assert "P1v3" not in plan["computeInventory"]
    assert matched["computeSku"] == "Standard_D2s_v5"
    assert unmatched["computeSku"] is None
    assert app["subscriptionId"] == "other"
    assert len(chart_records(report)) == 1
    path = tmp_path / f"inventory.{kind}"
    write_report(report, path, kind)
    text = reader(path)
    for expected in ("Standard_D2s_v5", "Windows_Server", "P1v3", "inventory-001", "not attributed"):
        assert expected in text


@pytest.mark.parametrize("kind,exporter,reader", [
    ("pdf", export_pdf, text_from_pdf), ("pptx", export_powerpoint, text_from_powerpoint),
])
def test_executive_exports_include_values_sources_evidence_and_partial_coverage(tmp_path, connected_result, kind, exporter, reader):
    report = build_report(connected_result)
    path = tmp_path / f"report.{kind}"
    exporter(report, path)
    text = reader(path)
    for expected in ["Savings plans", "4.25", "123.45", "evidence-001", "partial", "incomplete",
                     "2026-01-01", "synthetic-subscription", "not additive", "alternative"]:
        assert expected in text
    assert path.stat().st_size > 1000


def test_csv_schema_is_explicit_and_contains_all_records_with_provenance(tmp_path, connected_result):
    report = build_report(connected_result)
    path = tmp_path / "report.csv"
    export_csv(report, path)
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig", newline="")))
    assert len(rows) == 2
    assert rows[0]["hourlyCommitment"] == "4.25"
    assert rows[0]["savingsAmount"] == "123.45"
    assert rows[1]["modeledCost"] == ""
    assert all("incomplete" in row["sourceCoverage"] for row in rows)
    assert all(row["collectedAt"] == report.collected_at for row in rows)


def test_mixed_currencies_are_not_aggregated(connected_result):
    second = copy.deepcopy(connected_result["recommendations"][0])
    second["currencyCode"] = "EUR"
    connected_result["recommendations"].append(second)
    report = build_report(connected_result)
    assert report.metrics[-1] == ("Currencies (not summed)", "EUR, USD")
    assert {row["currency"] for row in report.records} == {"EUR", "USD"}
    assert not any("total savings" in label.lower() for label, _ in report.metrics)
    charts = report_charts(report)
    assert len(charts) == 3
    assert charts[0][1] == "savingsPercent"
    assert [chart[2] for chart in charts[1:]] == [" EUR/h", " USD/h"]
    assert all(len(rows) == 1 for _, _, _, rows in charts[1:])
    assert all(row["recordType"] == "savingsPlan" for _, _, _, rows in charts for row in rows)


def test_charts_do_not_invent_missing_amounts(connected_result):
    connected_result["recommendations"][0]["recommendationDetails"] = {}
    assert report_charts(build_report(connected_result)) == []


@pytest.mark.parametrize("kind,exporter,reader", [
    ("pdf", export_pdf, text_from_pdf), ("pptx", export_powerpoint, text_from_powerpoint),
    ("csv", export_csv, lambda path: path.read_text(encoding="utf-8-sig")),
])
def test_empty_denied_result_still_exports_coverage_not_success(tmp_path, connected_result, kind, exporter, reader):
    connected_result["recommendations"] = []
    report = build_report(connected_result)
    path = tmp_path / f"empty.{kind}"
    exporter(report, path)
    text = reader(path)
    assert "No recommendation records" in text
    assert "incomplete" in text
    assert "access denied" not in text
    assert "Request Reader access" not in text
    assert "4.25" not in text


def test_csv_formula_neutralization_and_shared_secret_redaction(tmp_path, connected_result):
    connected_result["advisorRecommendations"] = [{
        "recommendationReference": "advisor-001", "problem": "=HYPERLINK(\"bad\")",
        "solution": "Authorization: Bearer synthetic-private-token password=private-password",
        "impact": "High",
    }]
    report = build_report(connected_result)
    advisor = report.records[-1]
    assert "synthetic-private-token" not in advisor["notes"]
    assert "private-password" not in advisor["notes"]
    assert "[REDACTED]" in advisor["notes"]
    path = tmp_path / "safe.csv"
    export_csv(report, path)
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig")))
    assert rows[-1]["label"].startswith("'=")
    assert redacted_result({"access_token": "private", "nested": ["Bearer private"]}) == {
        "access_token": "[REDACTED]", "nested": ["[REDACTED]"],
    }


@pytest.mark.parametrize("kind,reader", [
    ("pdf", text_from_pdf), ("pptx", text_from_powerpoint),
    ("csv", lambda path: path.read_text(encoding="utf-8-sig")),
])
def test_secrets_are_redacted_in_every_report_format(tmp_path, connected_result, kind, reader):
    connected_result["advisorRecommendations"] = [{
        "problem": "Review workload", "solution": "Bearer secret-test-assertion client_secret=secret-test-value",
        "recommendationReference": "advisor-redacted", "impact": "High",
    }]
    path = tmp_path / f"safe.{kind}"
    write_report(build_report(connected_result), path, kind)
    text = reader(path)
    assert "secret-test-assertion" not in text
    assert "secret-test-value" not in text
    assert "[REDACTED]" in text


def test_failed_export_preserves_existing_file_and_removes_temporary_file(tmp_path, connected_result, monkeypatch):
    from api import reporting

    path = tmp_path / "existing.pdf"
    path.write_bytes(b"previous report")
    def fail(_report, destination):
        destination.write_bytes(b"incomplete new report")
        raise OSError("Disk write failed")
    monkeypatch.setattr(reporting, "export_pdf", fail)
    with pytest.raises(OSError, match="Disk write failed"):
        write_report(build_report(connected_result), path, "pdf")
    assert path.read_bytes() == b"previous report"
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize("kind,exporter,reader", [
    ("pdf", export_pdf, text_from_pdf), ("pptx", export_powerpoint, text_from_powerpoint),
])
def test_long_reports_paginate_without_losing_last_record_or_assumptions(tmp_path, connected_result, kind, exporter, reader):
    connected_result["advisorRecommendations"] = [{
        "recommendationReference": f"advisor-{index}", "problem": "Cost optimization & <review>",
        "solution": "Long recommendation text " * 100 + f"FINAL-{index}",
        "impact": "High", "lastUpdated": "2026-01-01",
    } for index in range(12)]
    path = tmp_path / f"long.{kind}"
    exporter(build_report(connected_result), path)
    text = reader(path)
    assert "FINAL-11" in text
    assert "advisor-11" in text
    if kind == "pdf":
        assert len(PdfReader(path).pages) > 2
    else:
        presentation = Presentation(path)
        assert len(presentation.slides) > 12
        for slide in presentation.slides:
            for shape in slide.shapes:
                assert shape.top + shape.height <= presentation.slide_height


@pytest.mark.parametrize("mode", ["what-if", "comparison"])
def test_offline_report_labels_estimates_and_preserves_option_values(mode):
    sample = Path(__file__).parents[1] / "samples" / "hourly-costs.csv"
    assumptions = {
        "mode": mode, "currencyCode": "USD", "expectedChangePercent": 0,
        **({"term": "P1Y", "discountPercent": 20, "scenario": "balanced"} if mode == "what-if"
           else {"discountPercentP1Y": 15, "discountPercentP3Y": 25}),
    }
    result = analyze_csv_file(sample, assumptions)
    report = build_report(result)
    assert len(report.records) == (1 if mode == "what-if" else 3)
    assert all(row["recordType"] == "model" for row in report.records)
    assert any("Offline planning estimate" in notice for notice in report.notices)
    assert "no Azure" in report.coverage[0]
    assert report.records[0]["savingsAmount"] == (
        result["netSavingsObserved"] if mode == "what-if" else result["options"][0]["netSavingsObserved"]
    )
